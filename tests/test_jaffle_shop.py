"""The bundled Jaffle Shop stays runnable and preserves semantics after rewriting."""

from collections import Counter
from decimal import Decimal
import io
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlglot import exp

from slbt.catalog import import_catalog
from slbt.catalog.export import catalog_outputs
from slbt.export import export_dbt
from slbt.inspection import expanded_query, inspect_query
from slbt.optimize import optimize_project
from slbt.project import Project, transformation_tree
from slbt.schema import load_yaml
from slbt.sql import Compiler


EXAMPLE = Path(__file__).resolve().parents[1] / "examples/jaffle_shop"
OPTIMIZATION_MODELS = (
    "gld_optimizer_constants", "gld_customer_lifetime_value",
    "gld_distinct_paying_customers", "gld_customer_order_ranking",
    "gld_high_value_customers", "gld_payment_union", "gld_customer_leaderboard",
)
OFFLINE_MODELS = (
    "brz_customers", "brz_orders", "brz_payments",
    "slv_customers", "slv_orders", "slv_payments",
    "gld_customers", "gld_orders", "gld_daily_sales", "gld_payment_method_summary",
    "gld_customer_cohorts", "gld_join_counts", "gld_macro_arguments",
    *OPTIMIZATION_MODELS,
)


@pytest.fixture
def loaded():
    return Project(EXAMPLE / "slbt_project.yaml")


@pytest.fixture
def connection():
    duckdb = pytest.importorskip("duckdb")
    with duckdb.connect() as connection:
        for name, count in (("customers", 100), ("orders", 99), ("payments", 113)):
            path = EXAMPLE / "dbt/seeds" / f"raw_{name}.csv"
            connection.execute(f"CREATE TABLE raw_{name} AS SELECT * FROM read_csv_auto(?)", [str(path)])
            assert connection.execute(f"SELECT COUNT(*) FROM raw_{name}").fetchone() == (count,)
        yield connection


def bound_sql(tree):
    """Bind logical source leaves only in the test execution harness."""
    tree = tree.copy()
    for table in tree.find_all(exp.Table):
        if table.catalog == "source" and table.db == "jaffle_shop":
            table.set("catalog", None)
            table.set("db", None)
            table.set("this", exp.to_identifier("raw_" + table.name))
    return tree.sql("duckdb")


def query(connection, tree):
    return connection.execute(bound_sql(tree)).fetchall()


@pytest.mark.parametrize("name", OFFLINE_MODELS)
def test_seeded_inspection_preserves_whole_and_selected_outputs(loaded, connection, name):
    complete = expanded_query(loaded, name)
    assert Counter(query(connection, inspect_query(loaded, name))) == Counter(query(connection, complete))
    column = complete.selects[-1].output_name
    expected = exp.select(exp.column(column, table="target")).from_(complete.subquery("target"))
    assert Counter(query(connection, inspect_query(loaded, name, [column]))) == Counter(query(connection, expected))


def test_marts_match_independent_jaffle_shop_calculations(loaded, connection):
    connection.execute("CREATE TEMP TABLE orders AS " + bound_sql(expanded_query(loaded, "gld_orders")))
    connection.execute("CREATE TEMP TABLE customers AS " + bound_sql(expanded_query(loaded, "gld_customers")))
    assert connection.execute("SELECT COUNT(*), SUM(amount) FROM orders").fetchone() == (99, Decimal("1672.00"))
    assert connection.execute("SELECT amount, payment_count FROM orders WHERE order_id = 25").fetchone() == (Decimal("58.00"), 3)
    assert connection.execute("SELECT status, amount FROM orders WHERE order_id = 1").fetchone() == ("returned", Decimal("10.00"))
    assert connection.execute("SELECT COUNT(*), SUM(customer_lifetime_value), SUM(number_of_orders) FROM customers").fetchone() == (100, Decimal("1672.00"), 99)
    assert connection.execute("SELECT COUNT(*) FROM customers WHERE number_of_orders = 0 AND customer_lifetime_value = 0 AND first_order IS NULL").fetchone() == (38,)
    expected = connection.execute("""
        SELECT c.id, COUNT(DISTINCT o.id), CAST(COALESCE(SUM(p.amount), 0) / 100.0 AS DECIMAL(12, 2))
        FROM raw_customers c
        LEFT JOIN raw_orders o ON c.id = o.user_id
        LEFT JOIN raw_payments p ON o.id = p.order_id
        GROUP BY c.id
    """).fetchall()
    actual = connection.execute("SELECT customer_id, number_of_orders, customer_lifetime_value FROM customers").fetchall()
    assert Counter(actual) == Counter(expected)


def test_gold_optimizer_rewrites_preserve_results_and_metadata(examples, loaded, connection):
    destination = examples
    copied = Project(destination / "slbt_project.yaml")
    result = optimize_project(copied, models=list(OPTIMIZATION_MODELS), preserve_marco=False)
    assert len(result.changed) == 7
    rewritten = Project(copied.path)
    originals = {model.name: model for model in loaded.models}
    for model in rewritten.models:
        if model.name not in OPTIMIZATION_MODELS:
            continue
        original = originals[model.name]
        assert model.authored["layer"] == "gold"
        assert Counter(query(connection, expanded_query(rewritten, model.name))) == Counter(query(connection, expanded_query(loaded, model.name)))
        assert list(model.output["columns"]) == list(original.output["columns"])
        for name, metadata in model.output["columns"].items():
            assert {key: value for key, value in metadata.items() if key != "expression"} == {key: value for key, value in original.output["columns"][name].items() if key != "expression"}
        assert model.config == original.config
    lifetime = next(model for model in rewritten.models if model.name == "gld_customer_lifetime_value")
    assert not lifetime.steps


@pytest.mark.parametrize("preserve_macro", [True, False])
def test_macro_arguments_optimization_preserves_results(examples, loaded, connection, preserve_macro):
    destination = examples
    copied = Project(destination / "slbt_project.yaml")
    result = optimize_project(copied, models=["gld_macro_arguments"], preserve_marco=preserve_macro)
    assert len(result.changed) == 1
    rewritten = Project(copied.path)
    model = next(model for model in rewritten.models if model.name == "gld_macro_arguments")
    tree = transformation_tree(model.output, Compiler(rewritten.macros, preserve_macro=True), model.name)
    assert any(call.name.lower() == "rounded_total" for call in tree.find_all(exp.Anonymous)) == preserve_macro
    original = next(model for model in loaded.models if model.name == "gld_macro_arguments")
    assert list(model.output["columns"]) == list(original.output["columns"])
    assert model.output["columns"]["rounded_with_fee"]["description"] == original.output["columns"]["rounded_with_fee"]["description"]
    assert Counter(query(connection, expanded_query(rewritten, model.name))) == Counter(query(connection, expanded_query(loaded, model.name)))


@pytest.mark.parametrize("preserve_macro", [True, False])
def test_full_project_optimization_skips_runtime_and_preserves_results(examples, loaded, connection, preserve_macro):
    destination = examples
    copied = Project(destination / "slbt_project.yaml")
    before = {path: path.read_bytes() for path in (destination / "models").rglob("*.yml")}
    preview = optimize_project(copied, dry_run=True, preserve_marco=preserve_macro)
    assert preview.skipped == {"gld_customer_attribute_hash": "macro 'hash_relation_attributes' requires dbt runtime and cannot be expanded offline"}
    assert len(preview.changed) == 14
    untouched = [destination / "models" / name for name in ("brz/customers.yml", "brz/orders.yml", "brz/payments.yml", "slv/customers.yml", "slv/orders.yml", "gld/payment_method_summary.yml")]
    assert not set(untouched) & set(preview.changed)
    assert set(preview.graphs) == set(OFFLINE_MODELS)
    assert {path: path.read_bytes() for path in before} == before
    result = optimize_project(copied, preserve_marco=preserve_macro)
    assert result.skipped == preview.skipped
    assert result.changed == preview.changed
    assert all(path.read_bytes() == before[path] for path in untouched)
    rewritten = Project(copied.path)
    originals = {model.name: model for model in loaded.models}
    for model in rewritten.models:
        original = originals[model.name]
        if model.name == "gld_customer_attribute_hash":
            assert model.authored == original.authored
            continue
        assert Counter(query(connection, expanded_query(rewritten, model.name))) == Counter(query(connection, expanded_query(loaded, model.name)))
        assert list(model.output["columns"]) == list(original.output["columns"])
        for name, metadata in model.output["columns"].items():
            assert {key: value for key, value in metadata.items() if key != "expression"} == {key: value for key, value in original.output["columns"][name].items() if key != "expression"}
    daily_sales = next(model for model in rewritten.models if model.name == "gld_daily_sales")
    tree = transformation_tree(daily_sales.output, Compiler(rewritten.macros, preserve_macro=True), daily_sales.name)
    assert any(call.name.lower() == "round_money" for call in tree.find_all(exp.Anonymous)) == preserve_macro


def test_export_populates_runnable_dbt_project_and_prunes_only_generated_files(examples):
    project = Project(examples / "slbt_project.yaml")
    destination = examples / "dbt"
    assert project.dbt_root == destination
    config = load_yaml(destination / "dbt_project.yml")
    for category in ("model", "macro", "snapshot", "seed", "test"):
        assert config[f"{category}-paths"] == [category + "s"]
    assert config["profile"] in load_yaml(destination / "profiles.yml")
    tracked = {path: path.read_bytes() for path in destination.rglob("*") if path.is_file()}
    assert {path.relative_to(destination).as_posix() for path in tracked} == {
        "dbt_project.yml", "profiles.yml", "seeds/LICENSE", "seeds/README.md",
        "seeds/raw_customers.csv", "seeds/raw_orders.csv", "seeds/raw_payments.csv",
        "tests/jaffle_shop_reconciliation.sql",
    }
    preview = export_dbt(project, dry_run=True)
    assert preview
    assert {path: path.read_bytes() for path in destination.rglob("*") if path.is_file()} == tracked
    assert export_dbt(project) == preview
    assert (destination / "models/gld/gld_customers.sql").is_file()
    assert any((destination / "macros").rglob("*.sql"))
    assert (destination / "snapshots/slv/customer_history.sql").is_file()
    assert not list(destination.glob("*.duckdb"))
    stale = []
    for category in ("models", "macros", "snapshots"):
        path = destination / category / "stale.sql"
        path.write_text("-- obsolete generated file\n", encoding="utf-8")
        stale.append(path)
    expected = [("remove", path) for path in sorted(stale)]
    assert export_dbt(project, prune=True, dry_run=True) == expected
    assert all(path.exists() for path in stale)
    assert export_dbt(project, prune=True) == expected
    assert all(not path.exists() for path in stale)
    assert {path: path.read_bytes() for path in tracked} == tracked
    assert export_dbt(project) == []


def test_bundled_catalog_is_current_formatted_and_imports_as_noop(loaded):
    path = EXAMPLE / "catalog/catalog.xlsx"
    actual = load_workbook(path)
    expected = load_workbook(io.BytesIO(catalog_outputs(loaded, "catalog")["catalog.xlsx"]))
    try:
        assert actual.sheetnames == expected.sheetnames
        assert set(actual.sheetnames) == {"sources", "brz", "slv", "gld", "macros", "defaults", "classifications"}
        for sheet in actual:
            assert list(sheet.values) == list(expected[sheet.title].values)
            assert sheet.freeze_panes == "C2"
            assert sheet.auto_filter.ref
            for row in sheet:
                for cell in row:
                    assert cell.data_type not in {"f", "e"}
                    assert cell.font.name == "Arial"
        assert import_catalog(loaded, path, dry_run=True) == []
    finally:
        actual.close()
        expected.close()
