"""Version-1 slbt contracts, including offline, whole-project mutation preflight.

The former single-model-document, model_defaults, argparse and eager-prune
expectations are deliberately replaced rather than preserved as compatibility APIs.
"""

import copy
import csv
import io
import json
from itertools import islice
from pathlib import Path, PurePosixPath

import pytest
import yaml
from click.testing import CliRunner
from openpyxl import load_workbook
from sqlglot import exp

from slbt.catalog import export_catalog, import_catalog
from slbt.catalog.export import catalog_outputs
from slbt.catalog.review import CLASSIFICATION_HEADERS, DEFAULT_HEADERS, HEADERS, MACRO_HEADERS, SOURCE_HEADERS, choose, model_rows, parse_json, workbook_bytes
from slbt.cli import cli
from slbt.compiler.dbt import compile_sql, dbt_outputs, schema
from slbt.definitions import classifications, defaults
from slbt.definitions.model import validate
from slbt.enrich import enrich_project
from slbt.export import export_dbt, reconcile
from slbt.optimize import _alphabetic_names, _name_generated_relations, optimize_project
from slbt.project import Project, load_project, resolve_project, transformation_tree
from slbt.scaffold import initialize
from slbt.schema import SpecError, dump_yaml, load_yaml
from slbt.sql import Compiler, read_query


@pytest.fixture
def root(tmp_path):
    initialize(tmp_path / "project")
    return tmp_path / "project"


def project(root):
    return Project(root / "slbt_project.yaml")


def edit(path, **updates):
    data = load_yaml(path)
    data.update(updates)
    path.write_bytes(dump_yaml(data))
    return data


def model(name="example", **updates):
    return {"name": name, "output": {"sql": "SELECT 1 AS id", "columns": {"id": {"data_type": "bigint"}}}, **updates}


def resource(root, tables=None, sources=None, name="example.yml", **updates):
    data = {"version": 1, "layer": "silver", "domain": "example", **updates}
    if tables is not None:
        data["tables"] = tables
    if sources is not None:
        data["sources"] = sources
    path = root / "models" / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(dump_yaml(data))
    return path


def baseline(root):
    return {p.relative_to(root): p.read_bytes() for p in root.rglob("*") if p.is_file()}


def book(root):
    export_catalog(project(root))
    return root / "catalog/catalog.xlsx"


def workbook_edit(path, title, kind, name, header, value):
    workbook = load_workbook(path)
    sheet = workbook[title]
    headers = [c.value for c in sheet[1]]
    for row in sheet.iter_rows(min_row=2):
        if row[0].value == kind and row[1].value == name:
            row[headers.index(header)].value = value
            break
    else:
        raise AssertionError((kind, name))
    workbook.save(path)
    workbook.close()


def source_group(**updates):
    return {"name": "raw", "schema": "landing", "tables": [{"name": "orders", "columns": [{"name": "id", "data_type": "bigint", "source_data_type": "BIGINT", "description": "Upstream identifier"}]}], **updates}


def source_catalog(root):
    config = load_yaml(root / "slbt_project.yaml")
    config["catalog"][0]["sheets"].append({"name": "Sources", "type": "source"})
    (root / "slbt_project.yaml").write_bytes(dump_yaml(config))


def test_bundled_project_pure_rendering(examples):
    before = baseline(examples)
    loaded = project(examples)
    assert len(loaded.models) == 21
    assert len(loaded.source_paths) == 3
    files = dbt_outputs(loaded)
    assert all(isinstance(k, str) and isinstance(v, bytes) for k, v in files.items())
    assert b"ref('slv_customers')" in files["models/gld/gld_customers.sql"]
    assert b"ref('gld_orders')" in files["models/gld/optimization/gld_distinct_paying_customers.sql"]
    assert b"source('jaffle_shop', 'orders')" in files["models/brz/brz_orders.sql"]
    assert b"snapshot customer_history" in files["snapshots/slv/customer_history.sql"]
    assert yaml.safe_load(files["models/gld/gld.yml"])["version"] == 2
    gold = dbt_outputs(loaded, PurePosixPath("gld"))
    assert len([path for path in gold if path.endswith(".sql")]) == 15
    assert all(path.startswith("models/gld/") for path in gold)
    assert baseline(examples) == before


def test_low_level_interfaces_and_domain_optional():
    resolved = {"version": 1, "layer": "custom", **model()}
    assert list(validate(resolved)) == ["slbt_output"]
    assert "SELECT 1 AS id" in compile_sql(resolved)
    assert yaml.safe_load(schema({"example": resolved}))["models"][0]["name"] == "example"
    with pytest.raises(ValueError):
        validate({"version": 1, "tables": [model()]})


@pytest.mark.parametrize("value", [True, "1", 2, None, 1.0])
def test_document_version_strict(root, value):
    resource(root, [model()], version=value)
    with pytest.raises(ValueError, match="version"):
        project(root)


@pytest.mark.parametrize("contents", ["version: 1\nversion: 1\n", "1: value\n", "true: value\n"])
def test_yaml_duplicate_and_nonstring_keys(root, contents):
    path = root / "invalid.yml"
    path.write_text(contents)
    with pytest.raises(ValueError):
        load_yaml(path)


def test_yaml_on_is_text_and_literals_are_finite(root):
    path = root / "boolean.yml"
    path.write_text("on: yes\nvalue: TRUE\nother: false\n")
    assert load_yaml(path) == {"on": "yes", "value": True, "other": False}
    for value in ['{"x":NaN}', '{"x":1e999}', '{"x":1,"x":2}']:
        with pytest.raises(ValueError):
            parse_json(value, dict, "Config")


@pytest.mark.parametrize("change", [
    {"model_defaults": {}}, {"optimization": {"preserve_macro": True}},
    {"optimization": {"preserve_macros": True}}, {"optimization": {"format": "source"}},
    {"optimization": {"enrich": "false"}}, {"dbt": {"path": "dbt/", "default_configs": None}},
    {"model_path": "../models"}, {"model_path": "C:/models"}, {"macro_path": "macros\\"},
    {"catalog_path": "models/catalog"}, {"dbt": {"path": "catalog/dbt"}},
])
def test_project_configuration_rejects_invalid(root, change):
    edit(root / "slbt_project.yaml", **change)
    with pytest.raises(ValueError):
        project(root)


@pytest.mark.parametrize("path_type", [Path, str])
def test_output_sibling_path_and_absolute_override(root, path_type):
    edit(root / "slbt_project.yaml", dbt={"path": "../generated"})
    loaded = project(root)
    assert loaded.destination("dbt") == root.parent / "generated"
    assert loaded.destination("dbt", path_type(root.parent / "override")) == root.parent / "override"
    assert loaded.destination("catalog", path_type(root.parent / "catalog-override")) == root.parent / "catalog-override"
    with pytest.raises(ValueError):
        loaded.destination("dbt", path_type(root / "models/future"))


def test_documents_identity_inheritance_and_shared_writes(root):
    path = resource(root, [model(), model("sibling")], [source_group()])
    loaded = project(root)
    selected = loaded.authored_models["example"]
    selected["description"] = "Reviewed"
    changes = loaded.rewrite_models({"example": selected})
    candidate = yaml.safe_load(changes[path])
    assert candidate["tables"][0]["description"] == "Reviewed"
    assert "layer" not in candidate["tables"][0]
    assert candidate["tables"][1] == model("sibling")
    assert candidate["sources"] == [source_group()]
    assert {m.name for m in loaded.models} == {"example", "sibling"}


def test_document_noop_preserves_override_bytes(root):
    path = resource(root, [model()])
    proposed = path.read_bytes().replace(b"name: example", b"name: example\n  description: 'Reviewed'  # keep this quote and comment")
    loaded = load_project(root / "slbt_project.yaml", {path: proposed})
    assert loaded.rewrite_models(loaded.authored_models)[path] == proposed


def test_document_projection_reordering_is_not_a_noop(root):
    path = resource(root, [model(output={"from": {"ref": "items", "alias": "v"}, "columns": {"id": "v.id", "other": "v.other"}})])
    loaded = project(root)
    selected = loaded.authored_models["example"]
    selected["output"]["columns"] = dict(reversed(selected["output"]["columns"].items()))
    changes = loaded.rewrite_models({"example": selected})
    assert changes[path] != path.read_bytes()
    assert list(load_yaml(path, changes)["tables"][0]["output"]["columns"]) == ["other", "id"]


@pytest.mark.parametrize("tables,sources", [([], []), ([model(), model()], None), ([{"version": 1, **model()}], None), (None, [{"name": "raw", "tables": []}])])
def test_invalid_documents(root, tables, sources):
    resource(root, tables, sources)
    with pytest.raises(ValueError):
        project(root)


def test_registry_required_exact_and_arbitrary_layers(root):
    path = root / "models/classifications.yaml"
    registry = load_yaml(path)
    registry["layers"][0]["name"] = "custom"
    path.write_bytes(dump_yaml(registry))
    resource(root, [model()], layer="custom")
    assert project(root).models[0].authored["layer"] == "custom"
    resource(root, [model()], layer="Custom")
    with pytest.raises(ValueError):
        project(root)


def test_sync_registry_order_and_dry_run(root):
    path = root / "models/classifications.yaml"
    registry = load_yaml(path)
    registry["layers"].insert(0, {"name": "bronze", "description": "Inputs", "domains": [{"name": "legacy", "description": "Legacy"}]})
    path.write_bytes(dump_yaml(registry))
    before = baseline(root)
    result = CliRunner().invoke(cli, ["--project-dir", str(root), "sync-classifications", "--dry-run"])
    assert result.exit_code == 0, result.output
    assert baseline(root) == before
    synced = classifications.synchronize(registry)
    assert [d["name"] for d in synced["layers"][0]["domains"]] == ["example", "legacy"]


def test_named_defaults_equal_overlap_and_nested_explicit_resolution(root):
    rules = {"configs": [
        {"name": "a", "target": {}, "config": {"meta": {"owner": "team", "a": 1}}},
        {"name": "b", "target": [{"layer": "silver"}, {"tables": "regex:^example$"}], "config": {"meta": {"owner": "team", "b": 2}}},
    ]}
    edit(root / "slbt_project.yaml", dbt={"path": "dbt/", "default_configs": rules})
    with pytest.raises(ValueError, match="overlapping.*a, b"):
        project(root)
    resource(root, [model(config={"meta": {"owner": "explicit"}})])
    assert project(root).models[0].config["meta"] == {"owner": "explicit", "a": 1, "b": 2, "domain": "example"}
    assert not defaults.equal(True, 1)


@pytest.mark.parametrize("target", [{"layer": []}, {"tables": []}, [], {"domain": "regex:"}, {"tables": "regex:["}, {"unknown": "x"}])
def test_defaults_target_contract(target):
    with pytest.raises(ValueError):
        defaults.validate({"configs": [{"name": "bad", "target": target}]})


@pytest.mark.parametrize("sql", ["DELETE FROM ref.a", "SELECT 1; SELECT 2", "SELECT 1 UNION SELECT 1, 2", "SELECT * FROM ref.a", "SELECT 1 AS other"])
def test_sql_query_and_projection_contract(root, sql):
    resource(root, [model(output={"sql": sql, "columns": {"id": {}}})])
    with pytest.raises(ValueError):
        project(root)


@pytest.mark.parametrize("output", [
    {"from": "ref.a", "columns": {"id": {"data_type": "bigint"}}},
    {"from": "ref.a", "sql": "SELECT 1 AS id", "columns": {"id": "1"}},
    {"from": "ref.a", "columns": {"id": "id"}, "limit": 0},
    {"from": "ref.a", "columns": {"id": "id"}, "limit": True},
    {"from": "ref.a", "columns": {"id": "id"}, "distinct": 1},
    {"from": "ref.a", "columns": {"id": "id"}, "order_by": []},
    {"from": "source.raw.orders", "joins": [{"ref": "a", "type": "left"}], "columns": {"id": "id"}},
    {"from": "ref.a", "joins": [{"ref": "b", "type": "cross", "on": "1=1"}], "columns": {"id": "id"}},
])
def test_structured_contract(root, output):
    resource(root, [model(output=output)])
    with pytest.raises(ValueError):
        project(root)


@pytest.mark.parametrize("steps,output", [
    ({"unused": {"sql": "SELECT 1 AS id"}}, {"sql": "SELECT 1 AS id", "columns": {"id": {}}}),
    ({}, {"from": "step.missing", "columns": {"id": "id"}}),
    ({"a": {"from": "step.b", "columns": {"id": "id"}}, "b": {"from": "step.a", "columns": {"id": "id"}}}, {"from": "step.a", "columns": {"id": "id"}}),
    ({"slbt_output": {"sql": "SELECT 1 AS id"}}, {"sql": "SELECT 1 AS id", "columns": {"id": {}}}),
])
def test_graph_contract(root, steps, output):
    resource(root, [model(steps=steps, output=output)])
    with pytest.raises(ValueError):
        project(root)


def test_local_sql_cte_does_not_capture_generated_step(root):
    resource(root, [model(steps={"a": {"sql": "SELECT 1 AS id"}}, output={"sql": "WITH __slbt_a AS (SELECT 2 AS id) SELECT s.id FROM step.a AS s JOIN __slbt_a AS c ON c.id > s.id", "columns": {"id": {}}})])
    result = dbt_outputs(project(root))["models/example.sql"].decode()
    assert "__slbt_a_2" in result
    assert "ref(" not in result


@pytest.mark.parametrize("config", [{"contract": {"enforced": "yes"}}, {"contract": {"enforced": True}}])
def test_contract_rejects_bad_enforcement_or_missing_type(root, config):
    resource(root, [model(config=config, output={"sql": "SELECT 1 AS id", "columns": {"id": {}}})])
    with pytest.raises(ValueError):
        project(root)


def test_snapshots_defaults_generation_and_collision(root):
    snapshot = {"name": "example_history", "config": {"strategy": "check", "unique_key": "id", "check_cols": "all"}, "columns": ["id"]}
    resource(root, [model(snapshot=snapshot)])
    rules = {"snapshot_configs": [{"name": "history", "target": {}, "config": {"target_schema": "history"}}]}
    edit(root / "slbt_project.yaml", dbt={"path": "dbt/", "default_configs": rules})
    loaded = project(root)
    assert loaded.models[0].snapshot["config"]["target_schema"] == "history"
    sql = dbt_outputs(loaded)["snapshots/example_history.sql"].decode()
    assert "snapshot example_history" in sql and "ref('example')" in sql
    assert "history" in sql
    resource(root, [model(snapshot={**snapshot, "name": "example"})])
    with pytest.raises(ValueError, match="collision"):
        project(root)


@pytest.mark.parametrize("snapshot", [
    {"name": "history", "columns": ["id"], "config": {"strategy": "timestamp", "unique_key": "id", "updated_at": "missing"}},
    {"name": "history", "columns": ["id"], "config": {"strategy": "check", "unique_key": "missing", "check_cols": "all"}},
    {"name": "history", "columns": [], "config": {}},
])
def test_invalid_snapshots(root, snapshot):
    resource(root, [model(snapshot=snapshot)])
    with pytest.raises(ValueError):
        project(root)


def test_source_registration_projection_and_shared_group_consistency(root):
    resource(root, [model()], [source_group()])
    loaded = project(root)
    table = loaded.sources["raw"]["tables"]["raw.orders"]
    assert table["columns"]["id"]["source_data_type"] == "BIGINT"
    generated = yaml.safe_load(dbt_outputs(loaded)["models/example.yml"])
    assert generated["sources"][0]["tables"][0]["config"]["meta"]["layer"] == "silver"
    assert generated["sources"][0]["tables"][0]["columns"][0]["config"]["meta"]["source_data_type"] == "BIGINT"
    resource(root, sources=[source_group(schema="different", tables=[{"name": "other"}])], name="other.yml")
    with pytest.raises(ValueError, match="conflicting source"):
        project(root)


@pytest.mark.parametrize("properties", [{"config": {"materialized": "table"}}, {"config": {"enabled": 1}}, {"freshness": {}}, {"description": ""}, {"quoting": []}])
def test_source_structural_rejections(root, properties):
    resource(root, [model()], [source_group(**properties)])
    with pytest.raises(ValueError):
        project(root)


@pytest.mark.parametrize("format", ["xlsx", "csv"])
def test_catalog_noop_roundtrip_and_metadata_edit(root, format):
    resource(root, [model(), model("sibling")], [source_group()])
    source_catalog(root)
    export_catalog(project(root), format=format)
    target = root / "catalog" / ("catalog.xlsx" if format == "xlsx" else "csv")
    before = baseline(root)
    assert import_catalog(project(root), target, dry_run=True) == []
    assert import_catalog(project(root), target) == []
    assert baseline(root) == before
    if format == "xlsx":
        workbook_edit(target, "Models", "TABLE", "example", "Description", "Reviewed")
    else:
        path = target / "Models.csv"
        rows = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8"))))
        rows[1][HEADERS.index("Description")] = "Reviewed"
        with path.open("w", encoding="utf-8", newline="") as stream:
            csv.writer(stream).writerows(rows)
    changes = import_catalog(project(root), target, dry_run=True)
    assert changes == [root / "models/example.yml"]
    assert "description" not in load_yaml(changes[0])["tables"][0]
    import_catalog(project(root), target)
    saved = load_yaml(changes[0])
    assert saved["tables"][0]["description"] == "Reviewed"
    assert saved["tables"][1] == model("sibling")
    assert "layer" not in saved["tables"][0]


def test_import_source_relation_metadata_preserves_authored_grammar(root):
    resource(root, [model(output={"from": {"source": "raw.orders", "alias": "o"}, "columns": {"id": "o.id"}})], [source_group()])
    target = book(root)
    workbook_edit(target, "Models", "COLUMN", "id", "Description", "Reviewed")
    import_catalog(project(root), target)
    assert load_yaml(root / "models/example.yml")["tables"][0]["output"]["from"]["source"] == "raw.orders"


@pytest.mark.parametrize("value", ["=1+1", 1, True, "#DIV/0!"])
def test_workbook_rejects_formula_error_and_native_values(root, value):
    path = book(root)
    workbook_edit(path, "Models", "TABLE", "example", "Description", value)
    before = baseline(root)
    with pytest.raises(ValueError):
        import_catalog(project(root), path)
    assert baseline(root) == before


@pytest.mark.parametrize("headers,kinds", [
    (HEADERS, ["TABLE", "STEP", "OUTPUT", "COLUMN"]),
    (SOURCE_HEADERS, ["SOURCE", "TABLE", "COLUMN"]),
    (MACRO_HEADERS, ["macros/"]),
    (DEFAULT_HEADERS, ["defaults"]),
    (CLASSIFICATION_HEADERS, ["gold"]),
])
def test_workbook_single_line_rows_use_15_point_height(headers, kinds):
    rows = [[kind, "example", *([""] * (len(headers) - 2))] for kind in kinds]
    workbook = load_workbook(io.BytesIO(workbook_bytes({"Review": (headers, rows)})))
    sheet = workbook["Review"]
    assert [sheet.row_dimensions[index].height for index in range(1, len(rows) + 2)] == [15] * (len(rows) + 1)
    workbook.close()


@pytest.mark.parametrize("expression,height", [
    ("first\nsecond", 30),
    ("x" * 79, 30),
    ("\n".join(["line"] * 11), 150),
])
def test_workbook_row_height_expands_for_wrapped_text(expression, height):
    row = ["COLUMN", "example", *([""] * (len(HEADERS) - 2))]
    row[HEADERS.index("Expression")] = expression
    workbook = load_workbook(io.BytesIO(workbook_bytes({"Review": (HEADERS, [row])})))
    assert workbook["Review"].row_dimensions[2].height == height
    workbook.close()


def test_workbook_only_contains_configured_editable_sheets(root):
    path = book(root)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Models", "classifications"]
    assert all(sheet.sheet_state == "visible" for sheet in workbook)
    assert all(not sheet.data_validations.dataValidation for sheet in workbook)
    workbook.create_sheet("Unexpected")
    workbook.save(path)
    workbook.close()
    before = baseline(root)
    with pytest.raises(ValueError, match="workbook sheets do not match"):
        import_catalog(project(root), path)
    assert baseline(root) == before


@pytest.mark.parametrize("name", ["Column Mapping", "_slbt_mapping"])
def test_former_viewer_sheet_names_are_available_for_editable_sheets(root, name):
    config = load_yaml(root / "slbt_project.yaml")
    config["catalog"][0]["sheets"][0]["name"] = name
    (root / "slbt_project.yaml").write_bytes(dump_yaml(config))
    path = book(root)
    workbook = load_workbook(path)
    assert name in workbook.sheetnames
    workbook.close()
    assert import_catalog(project(root), path) == []


def test_import_required_identity_inapplicable_and_key_summary(root):
    resource(root, [model()])
    path = book(root)
    workbook_edit(path, "Models", "COLUMN", "id", "Key", "P")
    before = baseline(root)
    with pytest.raises(ValueError, match="Key/References"):
        import_catalog(project(root), path)
    assert baseline(root) == before
    workbook_edit(path, "Models", "COLUMN", "id", "Key", "")
    workbook_edit(path, "Models", "OUTPUT", "example", "Layer", "silver")
    with pytest.raises(ValueError, match="inapplicable"):
        import_catalog(project(root), path)


def test_defaults_review_keeps_inherited_config_authored(root):
    rules = {"configs": [{"name": "all", "target": {}, "config": {"materialized": "view", "meta": {"owner": "team"}}}]}
    edit(root / "slbt_project.yaml", dbt={"path": "dbt/", "default_configs": rules})
    path = book(root)
    workbook_edit(path, "Models", "TABLE", "example", "Description", "Reviewed")
    import_catalog(project(root), path)
    assert "config" not in load_yaml(root / "models/example.yml")["tables"][0]


def test_source_group_review_updates_unselected_siblings(root):
    resource(root, [model()], [source_group()])
    resource(root, sources=[source_group(tables=[{"name": "unselected"}])], name="sibling.yml")
    source_catalog(root)
    config = load_yaml(root / "slbt_project.yaml")
    config["catalog"][0]["sheets"][1]["includes"] = {"tables": "raw.orders"}
    (root / "slbt_project.yaml").write_bytes(dump_yaml(config))
    path = book(root)
    workbook_edit(path, "Sources", "SOURCE", "raw", "Schema", "reviewed")
    import_catalog(project(root), path)
    saved = load_yaml(root / "models/sibling.yml")
    assert saved["sources"][0]["schema"] == "reviewed"
    assert saved["sources"][0]["tables"] == [{"name": "unselected"}]


def test_macro_review_preserves_siblings_and_checks_unselected_models(root):
    macro = root / "macros/math.sql"
    macro.write_text("-- surrounding\n{% macro plus_one(value) %}{{ value }} + 1{% endmacro %}\n{% macro sibling() %}2{% endmacro %}\n")
    resource(root, [model(output={"sql": "SELECT plus_one(1) AS id", "columns": {"id": {}}})])
    config = load_yaml(root / "slbt_project.yaml")
    config["catalog"][0]["macros"] = [{"name": "Macros"}]
    (root / "slbt_project.yaml").write_bytes(dump_yaml(config))
    path = book(root)
    workbook = load_workbook(path)
    sheet = workbook["Macros"]
    row = next(r for r in sheet.iter_rows(min_row=2) if r[1].value == "plus_one")
    row[2].value = "{% macro plus_one(value) %}{{ value }} + 2{% endmacro %}"
    workbook.save(path)
    workbook.close()
    import_catalog(project(root), path)
    assert macro.read_text().startswith("-- surrounding")
    assert "macro sibling()" in macro.read_text()
    assert "+ 2" in macro.read_text()


def test_json_semantic_resolved_read_only_and_runtime_definitions(examples):
    loaded = project(examples)
    files = catalog_outputs(loaded, choose(loaded)[0], "json")
    catalog = json.loads(files["catalog.json"])
    assert catalog["format"] == "slbt.catalog"
    assert catalog["representation"] == "resolved"
    assert catalog["format_version"] == 1
    runtime = next(m for m in catalog["models"] if m["name"] == "gld_customer_attribute_hash")
    assert runtime["expansion"]["status"] != "expanded"
    assert catalog["macros"]
    target = examples / "review.json"
    target.write_bytes(files["catalog.json"])
    with pytest.raises(ValueError):
        import_catalog(loaded, target)


def test_json_column_expressions_metadata_and_direct_step_lineage(root):
    (root / "macros/math.sql").write_text("{% macro plus_one(value) %}{{ value }} + 1{% endmacro %}\n")
    resource(root, [model(steps={"prepared": {
        "from": {"source": "raw.orders", "alias": "o"},
        "columns": {"id": {"expression": "plus_one(o.id)", "description": "Adjusted ID"}},
    }}, output={
        "from": {"step": "prepared", "alias": "p"},
        "columns": {"id": {"expression": "p.id + p.id", "data_type": "bigint", "description": "Twice the ID", "data_tests": ["not_null"], "config": {"meta": {"owner": "team"}}}},
    })], [source_group()])
    loaded = project(root)
    before = baseline(root)
    files = catalog_outputs(loaded, "catalog", "json")
    assert files == catalog_outputs(loaded, "catalog", "json")
    assert baseline(root) == before
    exported = json.loads(files["catalog.json"])["models"][0]
    column = exported["output"]["columns"][0]
    assert column == {
        "name": "id", "expression": "p.id + p.id", "data_type": "bigint", "description": "Twice the ID", "data_tests": ["not_null"], "config": {"meta": {"owner": "team"}},
        "lineage": {"status": "resolved", "upstream_columns": [{"relation": "step.prepared", "column": "id"}]},
    }
    step_column = exported["steps"][0]["columns"][0]
    assert step_column["expression"] == "o.id + 1"
    assert step_column["description"] == "Adjusted ID"
    assert step_column["lineage"] == {"status": "resolved", "upstream_columns": [{"relation": "source.raw.orders", "column": "id"}]}


@pytest.mark.parametrize("query, expected", [
    ("SELECT id AS id FROM ref.upstream", [("ref.upstream", "id")]),
    ("SELECT ref.upstream.id AS id FROM ref.upstream", [("ref.upstream", "id")]),
    ("SELECT source.raw.orders.id AS id FROM source.raw.orders", [("source.raw.orders", "id")]),
    ("WITH c AS (SELECT u.id + 1 AS renamed FROM ref.upstream u) SELECT c.renamed * 2 AS id FROM c", [("ref.upstream", "id")]),
    ("SELECT c.renamed AS id FROM (SELECT u.id AS renamed FROM ref.upstream u) c", [("ref.upstream", "id")]),
    ("SELECT u.id AS id FROM ref.upstream u UNION ALL SELECT o.id AS different_name FROM source.raw.orders o", [("ref.upstream", "id"), ("source.raw.orders", "id")]),
    ("SELECT u.id + o.id + u.id AS id FROM ref.upstream u JOIN source.raw.orders o ON u.id = o.id", [("ref.upstream", "id"), ("source.raw.orders", "id")]),
    ("SELECT u.id AS id FROM ref.upstream u JOIN source.raw.orders o ON u.id = o.id WHERE o.id > 0", [("ref.upstream", "id")]),
    ("SELECT ROW_NUMBER() OVER (PARTITION BY u.id ORDER BY o.id) AS id FROM ref.upstream u JOIN source.raw.orders o ON u.id = o.id", [("ref.upstream", "id"), ("source.raw.orders", "id")]),
    ("SELECT (SELECT MAX(o.id) FROM source.raw.orders o) AS id FROM ref.upstream u", [("source.raw.orders", "id")]),
    ("SELECT COUNT(*) AS id FROM ref.upstream", []),
    ("SELECT 42 AS id", []),
])
def test_json_raw_sql_column_lineage(root, query, expected):
    resource(root, [model("upstream"), model(output={"sql": query, "columns": {"id": {"data_type": "bigint"}}})], [source_group()])
    catalog = json.loads(catalog_outputs(project(root), "catalog", "json")["catalog.json"])
    column = next(item for item in catalog["models"] if item["name"] == "example")["output"]["columns"][0]
    assert column["expression"]
    assert column["lineage"] == {"status": "resolved", "upstream_columns": [{"relation": relation, "column": name} for relation, name in expected]}


def test_json_ambiguous_column_lineage_is_partial(root):
    resource(root, [model("upstream"), model(output={
        "sql": "SELECT id AS id, u.id AS known FROM ref.upstream u JOIN source.raw.orders o ON u.id = o.id",
        "columns": {"id": {}, "known": {}},
    })], [source_group()])
    catalog = json.loads(catalog_outputs(project(root), "catalog", "json")["catalog.json"])
    ambiguous, known = next(item for item in catalog["models"] if item["name"] == "example")["output"]["columns"]
    assert ambiguous["lineage"] == {"status": "partial", "upstream_columns": [], "unresolved_columns": ["id"]}
    assert known["lineage"] == {"status": "resolved", "upstream_columns": [{"relation": "ref.upstream", "column": "id"}]}


def test_json_missing_schema_does_not_disambiguate_columns(root):
    resource(root, [model("upstream"), model(output={
        "sql": "SELECT id AS id FROM ref.upstream u JOIN ref.external e ON u.id = e.id",
        "columns": {"id": {}},
    })])
    catalog = json.loads(catalog_outputs(project(root), "catalog", "json")["catalog.json"])
    column = next(item for item in catalog["models"] if item["name"] == "example")["output"]["columns"][0]
    assert column["lineage"] == {"status": "partial", "upstream_columns": [], "unresolved_columns": ["id"]}


def test_json_lineage_analysis_failure_does_not_block_export(root, monkeypatch):
    from sqlglot.errors import OptimizeError

    def unsupported(*args, **kwargs):
        raise OptimizeError("unsupported offline analysis")

    monkeypatch.setattr("slbt.catalog.lineage.qualify", unsupported)
    catalog = json.loads(catalog_outputs(project(root), "catalog", "json")["catalog.json"])
    column = catalog["models"][0]["output"]["columns"][0]
    assert column["expression"] == "1"
    assert column["lineage"] == {"status": "partial", "upstream_columns": []}


@pytest.mark.parametrize("dialect, quote", [("duckdb", '"'), ("spark", "`"), ("snowflake", '"')])
def test_json_lineage_preserves_quoted_column_names(root, dialect, quote):
    edit(root / "slbt_project.yaml", optimization={"dialect": dialect})
    column = f"{quote}customer_id{quote}"
    resource(root, [model("upstream", output={"sql": f"SELECT 1 AS {column}", "columns": {"customer_id": {}}}), model(output={"sql": f"SELECT u.{column} AS id FROM ref.upstream u", "columns": {"id": {}}})])
    catalog = json.loads(catalog_outputs(project(root), "catalog", "json")["catalog.json"])
    exported = next(item for item in catalog["models"] if item["name"] == "example")["output"]["columns"][0]
    assert exported["lineage"] == {"status": "resolved", "upstream_columns": [{"relation": "ref.upstream", "column": "customer_id"}]}


def test_json_runtime_lineage_is_deferred_only_for_affected_columns(examples):
    catalog = json.loads(catalog_outputs(project(examples), "catalog", "json")["catalog.json"])
    exported = next(item for item in catalog["models"] if item["name"] == "gld_customer_attribute_hash")
    identifier, runtime = exported["output"]["columns"]
    assert identifier["lineage"] == {"status": "resolved", "upstream_columns": [{"relation": "ref.slv_customers", "column": "customer_id"}]}
    assert runtime["lineage"] == {"status": "deferred", "upstream_columns": []}
    assert "HASH_RELATION_ATTRIBUTES" in runtime["expression"]


def test_json_runtime_lineage_through_cte_keeps_known_dependencies(root):
    (root / "macros/runtime.sql").write_text("{% macro dynamic(value) %}{{ return(value) }}{% endmacro %}\n")
    resource(root, [model("upstream"), model(output={
        "sql": "WITH c AS (SELECT dynamic(u.id) AS opaque, u.id AS known FROM ref.upstream u) SELECT c.opaque + c.known AS id FROM c",
        "columns": {"id": {}},
    })])
    catalog = json.loads(catalog_outputs(project(root), "catalog", "json")["catalog.json"])
    column = next(item for item in catalog["models"] if item["name"] == "example")["output"]["columns"][0]
    assert column["lineage"] == {"status": "deferred", "upstream_columns": [{"relation": "ref.upstream", "column": "id"}]}


def test_json_raw_step_without_declared_columns_and_source_columns(root):
    resource(root, [model(steps={"prepared": {"sql": "SELECT o.id FROM source.raw.orders o"}}, output={
        "sql": "SELECT p.id AS id FROM step.prepared p", "columns": {"id": {}},
    })], [source_group()])
    source_catalog(root)
    catalog = json.loads(catalog_outputs(project(root), "catalog", "json")["catalog.json"])
    exported = catalog["models"][0]
    assert exported["steps"][0]["columns"] == []
    assert exported["output"]["columns"][0]["lineage"] == {"status": "resolved", "upstream_columns": [{"relation": "step.prepared", "column": "id"}]}
    source_column = catalog["sources"][0]["tables"][0]["columns"][0]
    assert source_column["name"] == "id"
    assert "expression" not in source_column
    assert "lineage" not in source_column


def test_json_lineage_retains_references_outside_selected_catalog(root):
    resource(root, [model("upstream"), model(output={"sql": "SELECT u.id AS id FROM ref.upstream u", "columns": {"id": {}}})])
    config = load_yaml(root / "slbt_project.yaml")
    config["catalog"][0]["sheets"][0]["includes"] = {"tables": "example"}
    (root / "slbt_project.yaml").write_bytes(dump_yaml(config))
    catalog = json.loads(catalog_outputs(project(root), "catalog", "json")["catalog.json"])
    assert len(catalog["models"]) == 1
    exported = catalog["models"][0]
    assert exported["dependencies"] == [{"relation": "ref.upstream", "included": False}]
    assert exported["output"]["columns"][0]["lineage"] == {"status": "resolved", "upstream_columns": [{"relation": "ref.upstream", "column": "id"}]}


@pytest.mark.parametrize("format", ["pretty", "compact", "source"])
def test_dbt_formats_separate_from_optimization(root, format):
    resource(root, [model(output={"sql": "select  1 as id -- layout\n", "columns": {"id": {}}})])
    edit(root / "slbt_project.yaml", dbt={"path": "dbt/", "format": format})
    text = dbt_outputs(project(root))["models/example.sql"].decode()
    if format == "source":
        assert "select  1 as id -- layout" in text
    else:
        assert "SELECT" in text


@pytest.mark.parametrize("format", ["pretty", "compact", "source"])
def test_dbt_export_formats_model_and_snapshot_configs(root, format):
    resource(root, [model(
        config={"materialized": "view", "meta": {"owner": "payments"}, "tags": ["payments", "sales"]},
        snapshot={"name": "example_history", "config": {
            "strategy": "check", "unique_key": "id", "check_cols": ["id"],
            "meta": {"owner": "payments"},
        }, "columns": ["id"]},
        output={"sql": "select  1 as id -- layout\n", "columns": {"id": {}}},
    )])
    edit(root / "slbt_project.yaml", dbt={"path": "dbt/", "format": format})
    files = dbt_outputs(project(root))
    model_sql = files["models/example.sql"].decode()
    snapshot_sql = files["snapshots/example_history.sql"].decode()
    for text in (model_sql, snapshot_sql):
        if format == "compact":
            assert "{{ config({" in text
            assert "'meta': {'owner': 'payments'" in text
        else:
            assert "{{\n  config({\n" in text
            assert "'meta': {\n      'owner': 'payments'" in text
    if format == "source":
        assert "select  1 as id -- layout\n" in model_sql


def test_reconciliation_unchanged_opt_in_prune_and_preview(root):
    loaded = project(root)
    before = baseline(root)
    assert export_dbt(loaded, dry_run=True)
    assert baseline(root) == before
    assert export_dbt(loaded)
    target = loaded.dbt_root / "models/example.sql"
    timestamp = target.stat().st_mtime_ns
    obsolete = loaded.dbt_root / "models/hand-authored.txt"
    obsolete.write_text("managed trees have no ownership manifest")
    assert export_dbt(project(root)) == []
    assert target.stat().st_mtime_ns == timestamp
    assert obsolete.exists()
    assert ("remove", obsolete) in export_dbt(project(root), prune=True, dry_run=True)
    assert obsolete.exists()
    export_dbt(project(root), prune=True)
    assert not obsolete.exists()


@pytest.mark.parametrize("path", ["../escape.sql", "/models/x.sql", "models/../x.sql", "models\\x.sql", "models/x:bad.sql", "outside/x.sql", "models/x.txt"])
def test_reconcile_rejects_unsafe_output_names(root, path):
    with pytest.raises(ValueError):
        reconcile(root / "dbt", "dbt", {path: b"SELECT 1"})
    assert not (root / "dbt").exists()


def test_reconcile_full_preflight_before_writes(root):
    destination = root / "dbt"
    (destination / "models").mkdir(parents=True)
    (destination / "models/block.sql").mkdir()
    with pytest.raises(ValueError):
        reconcile(destination, "dbt", {"models/first.sql": b"new", "models/block.sql": b"invalid"})
    assert not (destination / "models/first.sql").exists()


def test_named_catalog_pruning_suppressed_and_mixed_formats(root):
    edit(root / "slbt_project.yaml", catalog=[{"name": "catalog", "format": "json"}, {"name": "other", "format": "xlsx"}])
    export_catalog(project(root))
    assert (root / "catalog/catalog.json").exists()
    assert (root / "catalog/other.xlsx").exists()
    obsolete = root / "catalog/obsolete.json"
    obsolete.write_text("{}")
    assert export_catalog(project(root), names=["catalog"], prune=True) == []
    assert obsolete.exists()
    assert ("remove", obsolete) in export_catalog(project(root), prune=True)


def test_optimize_enrich_shared_document_preview_and_reload(root):
    resource(root, [model(output={"sql": "SELECT 1 + 2 AS id", "columns": {"id": {}}}), model("sibling")])
    before = baseline(root)
    result = optimize_project(project(root), models=["example"], enrich=True, dry_run=True)
    assert result.changed
    assert baseline(root) == before
    optimize_project(project(root), models=["example"], enrich=True)
    saved = load_yaml(root / "models/example.yml")
    assert saved["tables"][1] == model("sibling")
    assert "3" in saved["tables"][0]["output"]["sql"]
    assert saved["tables"][0]["output"]["columns"]["id"]["data_type"]
    path = root / "models/example.yml"
    before = baseline(root)
    timestamp = path.stat().st_mtime_ns
    assert optimize_project(project(root), models=["example"], enrich=True, dry_run=True).changed == []
    assert optimize_project(project(root), models=["example"], enrich=True).changed == []
    assert baseline(root) == before
    assert path.stat().st_mtime_ns == timestamp


@pytest.mark.parametrize("dry_run", [True, False])
def test_optimizer_noop_preserves_authored_formatting(root, dry_run):
    resource(root, sources=[source_group()], name="sources.yml")
    path = root / "models/example.yml"
    path.write_bytes(
        b"# Keep authored formatting and CRLF\r\nversion: 1\r\nlayer: silver\r\ndomain: example\r\ntables:\r\n"
        b"  - name: example\r\n    description: 'Already optimized' # keep this comment\r\n"
        b"    constraints:\r\n      - type: primary_key\r\n        columns:\r\n          - id\r\n"
        b"    output:\r\n      from:\r\n        source: raw.orders\r\n        alias: v\r\n"
        b"      columns:\r\n        id: v.id\r\n"
    )
    before = baseline(root)
    timestamp = path.stat().st_mtime_ns
    result = optimize_project(project(root), dry_run=dry_run)
    assert result.changed == []
    assert result.graphs
    assert baseline(root) == before
    assert path.stat().st_mtime_ns == timestamp


def test_optimizer_changed_model_uses_indented_lists_and_unquoted_on(root):
    path = resource(root, [model(constraints=[{"type": "primary_key", "columns": ["id"]}],
                                 output={"sql": "SELECT a.id + 0 AS id FROM ref.items AS a JOIN ref.other AS b ON a.id = b.id", "columns": {"id": {}}})])
    result = optimize_project(project(root))
    assert result.changed == [path]
    rendered = path.read_bytes()
    assert b"tables:\n  - name: example\n" in rendered
    assert b"    constraints:\n      - type: primary_key\n        columns:\n          - id\n" in rendered
    assert b"          on: a.id = b.id\n" in rendered
    assert b"'on':" not in rendered
    assert optimize_project(project(root)).changed == []


def test_alphabetic_names_skip_collisions_and_continue_past_z():
    assert list(islice(_alphabetic_names({"a", "C", "z", "aa"}), 26)) == [
        "b", *list("defghijklmnopqrstuvwxy"), "ab", "ac", "ad",
    ]
    names = list(islice(_alphabetic_names(set()), 704))
    assert names[:3] == ["a", "b", "c"]
    assert names[25:28] == ["z", "aa", "ab"]
    assert names[701:704] == ["zz", "aaa", "aab"]


@pytest.mark.parametrize("preserve_macro", [True, False])
def test_generated_step_and_alias_share_name_skipping_authored_steps(root, preserve_macro):
    path = resource(root, [model(
        steps={"a": {"sql": "SELECT candidate.id AS id FROM ref.items AS candidate", "columns": {"id": {}}}},
        output={"sql": "SELECT candidate.id AS id FROM step.a AS candidate "
                       "WHERE candidate.id > (SELECT AVG(benchmark.id) FROM ref.items AS benchmark)",
                "columns": {"id": {"description": "Keep the interface"}}},
    )])
    before = path.read_bytes()
    preview = optimize_project(project(root), preserve_marco=preserve_macro, dry_run=True)
    assert preview.changed == [path]
    assert path.read_bytes() == before
    optimize_project(project(root), preserve_marco=preserve_macro)
    saved = load_yaml(path)["tables"][0]
    assert list(saved["steps"]) == ["b"]
    assert saved["steps"]["b"]["from"]["alias"] == "benchmark"
    assert saved["output"]["from"]["alias"] == "candidate"
    assert saved["output"]["joins"] == [{"step": "b", "type": "inner", "on": "b._col_0 < candidate.id"}]
    assert saved["output"]["columns"]["id"]["description"] == "Keep the interface"
    assert optimize_project(project(root), preserve_marco=preserve_macro).changed == []


def test_generated_relation_alias_skips_authored_alias_and_updates_macro_argument(root):
    (root / "macros/math.sql").write_text("{% macro identity(value) %}{{ value }}{% endmacro %}")
    resource(root, [model(output={
        "sql": "SELECT identity(items.id) AS id FROM ref.items JOIN ref.other AS a ON items.id = a.id",
        "columns": {"id": {}},
    })])
    optimize_project(project(root), preserve_marco=True)
    saved = project(root).models[0]
    assert saved.output["from"] == {"ref": "items", "alias": "b"}
    assert saved.output["joins"] == [{"ref": "other", "alias": "a", "type": "inner", "on": "a.id = b.id"}]
    tree = transformation_tree(saved.output, Compiler(project(root).macros, preserve_macro=True), saved.name)
    call = next(tree.find_all(exp.Anonymous))
    assert call.expressions[0].sql() == "b.id"
    assert optimize_project(project(root)).changed == []


def test_generated_names_preserve_explicit_internal_looking_aliases_and_ctes(root):
    resource(root, [model(output={
        "sql": "WITH _u_0 AS (SELECT x.id FROM ref.items AS x GROUP BY x.id) "
               "SELECT _u_7.id AS id FROM _u_0 AS _u_7 "
               "WHERE _u_7.id > (SELECT AVG(benchmark.id) FROM ref.other AS benchmark)",
        "columns": {"id": {}},
    })])
    optimize_project(project(root))
    saved = project(root).models[0]
    assert "_u_0" in saved.steps
    assert saved.output["from"] == {"step": "_u_0", "alias": "_u_7"}
    assert "a" in saved.steps
    assert saved.output["joins"][0]["step"] == "a"
    assert "alias" not in saved.output["joins"][0]


@pytest.mark.parametrize("preserve_macro", [True, False])
def test_extracted_step_uses_authored_subquery_alias_without_redundant_alias(root, preserve_macro):
    import duckdb
    from jinja2 import Environment
    edit(root / "slbt_project.yaml", optimization={"dialect": "duckdb"})
    authored = (
        "SELECT paid.customer_id FROM ("
        "SELECT DISTINCT o.customer_id, o.order_id FROM ref.gld_orders AS o WHERE o.amount > 0"
        ") AS paid"
    )
    path = resource(root, [model(output={"sql": authored, "columns": {
        "customer_id": {"description": "One row per paid customer/order pair"},
    }})])
    before = path.read_bytes()
    preview = optimize_project(project(root), preserve_marco=preserve_macro, dry_run=True)
    assert preview.changed == [path]
    assert path.read_bytes() == before
    compiled_before = dbt_outputs(project(root))["models/example.sql"].decode()
    optimize_project(project(root), preserve_marco=preserve_macro)
    saved = project(root).models[0]
    assert list(saved.steps) == ["paid"]
    assert saved.steps["paid"]["from"] == {"ref": "gld_orders", "alias": "o"}
    assert saved.steps["paid"]["distinct"] is True
    assert saved.output["from"] == {"step": "paid"}
    assert saved.output["columns"]["customer_id"] == {
        "expression": "paid.customer_id", "description": "One row per paid customer/order pair",
    }
    compiled_after = dbt_outputs(project(root))["models/example.sql"].decode()
    with duckdb.connect() as connection:
        connection.execute("CREATE TABLE gld_orders AS SELECT * FROM (VALUES "
                           "(1, 10, 5), (1, 10, 5), (1, 11, 3), (2, 12, 0), (3, 13, 7)) "
                           "AS orders(customer_id, order_id, amount)")
        def render(value):
            return Environment().from_string(value).render(config=lambda value: "", ref=lambda name: name)

        expected = connection.execute(render(compiled_before)).fetchall()
        assert sorted(expected) == [(1,), (1,), (3,)]
        assert sorted(connection.execute(render(compiled_after)).fetchall()) == sorted(expected)
    assert optimize_project(project(root), preserve_marco=preserve_macro).changed == []


def test_extracted_named_steps_keep_names_in_nested_dependencies(root):
    resource(root, [model(output={
        "sql": "SELECT scored.id FROM (SELECT DISTINCT paid.id FROM "
               "(SELECT DISTINCT o.id, o.kind FROM ref.items AS o) AS paid) AS scored",
        "columns": {"id": {}},
    })])
    optimize_project(project(root))
    saved = project(root).models[0]
    assert list(saved.steps) == ["paid", "scored"]
    assert saved.steps["scored"]["from"] == {"step": "paid"}
    assert saved.output["from"] == {"step": "scored"}
    assert saved.output["columns"]["id"]["expression"] == "scored.id"
    assert optimize_project(project(root)).changed == []


def test_extracted_named_step_reserves_alias_before_alphabetic_fallback(root):
    resource(root, [model(output={
        "sql": "SELECT a.id FROM (SELECT DISTINCT o.id, o.kind FROM ref.items AS o) AS a "
               "WHERE a.id > (SELECT AVG(benchmark.id) FROM ref.other AS benchmark)",
        "columns": {"id": {}},
    })])
    optimize_project(project(root))
    saved = project(root).models[0]
    assert set(saved.steps) == {"a", "b"}
    assert saved.output["from"] == {"step": "a"}
    assert saved.output["joins"] == [{"step": "b", "type": "inner", "on": "a.id > b._col_0"}]
    assert optimize_project(project(root)).changed == []


def test_repeated_subquery_alias_uses_fallback_without_shadowing(root):
    resource(root, [model(output={
        "sql": "SELECT paid.id FROM (SELECT DISTINCT paid.id FROM "
               "(SELECT DISTINCT o.id, o.kind FROM ref.items AS o) AS paid) AS paid",
        "columns": {"id": {}},
    })])
    optimize_project(project(root))
    saved = project(root).models[0]
    assert list(saved.steps) == ["paid", "a"]
    assert saved.steps["a"]["from"] == {"step": "paid"}
    assert saved.output["from"] == {"step": "a", "alias": "paid"}
    assert saved.output["columns"]["id"]["expression"] == "paid.id"
    assert optimize_project(project(root)).changed == []


def test_extracted_alias_conflict_does_not_rename_authored_step(root):
    resource(root, [model(
        steps={"paid": {"sql": "SELECT DISTINCT x.id FROM ref.other AS x", "columns": {"id": {}},
                        "description": "Keep this authored step"}},
        output={"sql": "SELECT paid.id FROM (SELECT DISTINCT o.id, o.kind FROM ref.items AS o) AS paid "
                       "JOIN step.paid AS existing ON paid.id = existing.id", "columns": {"id": {}}},
    )])
    optimize_project(project(root))
    saved = project(root).models[0]
    assert set(saved.steps) == {"paid", "a"}
    assert saved.steps["paid"]["description"] == "Keep this authored step"
    assert saved.output["from"] == {"step": "a", "alias": "paid"}
    assert saved.output["joins"][0]["step"] == "paid"
    assert saved.output["joins"][0]["alias"] == "existing"
    assert optimize_project(project(root)).changed == []


def test_shared_extracted_step_keeps_distinct_self_join_alias(root):
    resource(root, [model(output={
        "sql": "SELECT paid.id FROM (SELECT DISTINCT o.id FROM ref.items AS o) AS paid "
               "JOIN (SELECT DISTINCT o.id FROM ref.items AS o) AS other ON paid.id = other.id",
        "columns": {"id": {}},
    })])
    optimize_project(project(root))
    saved = project(root).models[0]
    assert list(saved.steps) == ["paid"]
    assert saved.output["from"] == {"step": "paid"}
    assert saved.output["joins"] == [{"step": "paid", "alias": "other", "type": "inner", "on": "other.id = paid.id"}]
    assert optimize_project(project(root)).changed == []


def test_generated_alias_renaming_respects_correlations_and_local_shadowing(root):
    resource(root, [model()])
    saved = project(root).models[0]
    original = read_query("SELECT items.id FROM ref.items WHERE EXISTS "
                          "(SELECT 1 FROM ref.other AS a WHERE a.id = items.id)", "test", "duckdb")
    optimized = read_query("SELECT _q_0.id FROM ref.items AS _q_0 WHERE EXISTS "
                           "(SELECT 1 FROM ref.other AS a WHERE a.id = _q_0.id "
                           "AND EXISTS (SELECT 1 FROM ref.third AS _q_0 WHERE _q_0.id = a.id))", "test", "duckdb")
    _name_generated_relations(optimized, original, saved)
    assert optimized.sql("duckdb") == (
        'SELECT "c".id FROM ref.items AS "c" WHERE EXISTS(SELECT 1 FROM ref.other AS a '
        'WHERE a.id = "c".id AND EXISTS(SELECT 1 FROM ref.third AS "b" WHERE "b".id = a.id))'
    )


def test_generated_cte_and_alias_names_continue_past_z(root):
    import duckdb
    resource(root, [model()])
    saved = project(root).models[0]
    original = read_query("SELECT 0 AS total", "test", "duckdb")
    ctes = ", ".join(f"_u_{i} AS (SELECT {i} AS id)" for i in range(28))
    total = " + ".join(f"_u_{i}.id" for i in range(28))
    sources = " CROSS JOIN ".join(f"_u_{i} AS _u_{i}" for i in range(28))
    optimized = read_query(f"WITH {ctes} SELECT {total} AS total FROM {sources}", "test", "duckdb")
    before = optimized.sql("duckdb")
    _name_generated_relations(optimized, original, saved)
    names = list(islice(_alphabetic_names(set()), 28))
    assert [cte.alias for cte in optimized.find_all(exp.CTE)] == names
    tables = list(optimized.find_all(exp.Table))
    assert [table.name for table in tables] == names
    assert [table.alias for table in tables] == names
    assert {column.table for column in optimized.find_all(exp.Column)} == set(names)
    with duckdb.connect() as connection:
        assert connection.execute(before).fetchall() == connection.execute(optimized.sql("duckdb")).fetchall()


@pytest.mark.parametrize("materialized", ["", "MATERIALIZED "])
def test_generated_names_preserve_subquery_results_and_are_idempotent(root, materialized):
    import duckdb
    from jinja2 import Environment
    edit(root / "slbt_project.yaml", optimization={"dialect": "duckdb"})
    authored = (
        f"WITH a AS {materialized}(SELECT 1 AS id UNION ALL SELECT 2 AS id UNION ALL SELECT 3 AS id) "
        "SELECT a.id AS id, (SELECT AVG(benchmark.id) FROM a AS benchmark) AS average_id "
        "FROM a WHERE a.id > (SELECT AVG(benchmark.id) FROM a AS benchmark) "
        "AND EXISTS (SELECT 1 FROM (VALUES (1), (3), (3)) AS eligible(id) WHERE eligible.id = a.id)"
    )
    resource(root, [model(output={"sql": authored, "columns": {"id": {}, "average_id": {}}})])
    optimize_project(project(root))
    compiled = dbt_outputs(project(root))["models/example.sql"].decode()
    executable = Environment().from_string(compiled).render(config=lambda value: "")
    with duckdb.connect() as connection:
        assert connection.execute(authored).fetchall() == connection.execute(executable).fetchall()
    assert optimize_project(project(root)).changed == []


def test_enrichment_upstream_selected_only(root):
    resource(root, [model(output={"from": "source.raw.orders", "columns": {"id": "id"}}), model("downstream", output={"from": "ref.example", "columns": {"id": "id"}})], [source_group()])
    changed, report = enrich_project(project(root), ["downstream"])
    assert changed
    tables = load_yaml(root / "models/example.yml")["tables"]
    assert tables[0]["output"]["columns"]["id"] == "id"
    assert tables[1]["output"]["columns"]["id"]["description"] == "Upstream identifier"
    assert any(r.get("value") == "bigint" for r in report)


@pytest.mark.parametrize("preserve_macro", [True, False])
@pytest.mark.parametrize("enrich", [True, False])
def test_optimizer_skips_runtime_without_writes(root, preserve_macro, enrich):
    (root / "macros/runtime.sql").write_text("{% macro runtime(value) %}{{ adapter.quote(value) }}{% endmacro %}")
    resource(root, [model(output={"sql": "SELECT runtime(1) AS id", "columns": {"id": {}}})])
    before = baseline(root)
    result = optimize_project(project(root), preserve_marco=preserve_macro, enrich=enrich)
    assert result.skipped == {"example": "macro 'runtime' requires dbt runtime and cannot be expanded offline"}
    assert result.changed == []
    assert result.graphs == {}
    assert not getattr(result, "report", [])
    assert baseline(root) == before


@pytest.mark.parametrize("dry_run", [True, False])
def test_optimizer_runtime_skip_preserves_shared_document_model(root, dry_run):
    (root / "macros/runtime.sql").write_text("{% macro runtime(value) %}{{ adapter.quote(value) }}{% endmacro %}")
    skipped = model("runtime_model", output={"sql": "SELECT runtime(1) AS id", "columns": {"id": {}}})
    resource(root, [skipped, model(output={"sql": "SELECT 1 + 2 AS id", "columns": {"id": {}}})])
    before = baseline(root)
    result = optimize_project(project(root), dry_run=dry_run, enrich=True)
    assert set(result.skipped) == {"runtime_model"}
    assert set(result.graphs) == {"example"}
    assert len(result.changed) == 1
    assert all(item["model"] == "example" for item in result.report)
    if dry_run:
        assert baseline(root) == before
    else:
        saved = load_yaml(root / "models/example.yml")["tables"]
        assert saved[0] == skipped
        assert "3" in saved[1]["output"]["sql"]
        assert saved[1]["output"]["columns"]["id"]["data_type"]


def test_optimizer_runtime_skip_reports_wrappers_and_multiple_calls(root):
    (root / "macros/runtime.sql").write_text(
        "{% macro runtime(value) %}{{ adapter.quote(value) }}{% endmacro %}\n"
        "{% macro wrapper(value) %}{{ runtime(value) }}{% endmacro %}\n"
        "{% macro other(value) %}{{ run_query(value) }}{% endmacro %}"
    )
    resource(root, [model(output={"sql": "SELECT wrapper(1) + other(2) + wrapper(3) AS id", "columns": {"id": {}}})])
    result = optimize_project(project(root))
    assert result.skipped == {"example": "macros 'other', 'wrapper' require dbt runtime and cannot be expanded offline"}


@pytest.mark.parametrize("dry_run", [True, False])
def test_cli_optimizer_reports_runtime_skip(root, dry_run):
    (root / "macros/runtime.sql").write_text("{% macro runtime(value) %}{{ adapter.quote(value) }}{% endmacro %}")
    resource(root, [model(output={"sql": "SELECT runtime(1) AS id", "columns": {"id": {}}})])
    before = baseline(root)
    result = CliRunner().invoke(cli, ["--project-dir", str(root), "optimize", *(["--dry-run"] if dry_run else [])])
    assert result.exit_code == 0, result.output
    assert result.stderr == "Skipped example: macro 'runtime' requires dbt runtime and cannot be expanded offline. Model unchanged.\n"
    assert "0 file change(s)" in result.stdout
    assert baseline(root) == before


def test_optimizer_rejects_ambiguous_columns_without_writes(root):
    resource(root, [model(output={"sql": "SELECT id FROM ref.a AS a JOIN ref.b AS b ON a.id=b.id", "columns": {"id": {}}})])
    before = baseline(root)
    with pytest.raises(ValueError):
        optimize_project(project(root))
    assert baseline(root) == before


def test_optimizer_runtime_skip_does_not_hide_other_errors(root):
    (root / "macros/runtime.sql").write_text("{% macro runtime(value) %}{{ adapter.quote(value) }}{% endmacro %}")
    resource(root, [model("runtime_model", output={"sql": "SELECT runtime(1) AS id", "columns": {"id": {}}}),
                    model(output={"sql": "SELECT 1 + 2 AS id", "columns": {"id": {}}}),
                    model("invalid", output={"sql": "WITH RECURSIVE x AS (SELECT 1 AS id UNION ALL SELECT id + 1 FROM x WHERE id < 3) SELECT id FROM x", "columns": {"id": {}}})])
    before = baseline(root)
    with pytest.raises(SpecError, match="invalid: cannot optimize: recursive queries"):
        optimize_project(project(root))
    assert baseline(root) == before


@pytest.mark.parametrize("expression", [
    "passthrough(SUM(v.id))", "passthrough(SUM(plus_one(v.id)))",
    "passthrough(passthrough(SUM(v.id)))", "passthrough(value => SUM(v.id))",
    "passthrough(SUM(v.id) OVER ())",
])
def test_optimizer_preserves_macros_with_visible_query_semantics(root, expression):
    (root / "macros/math.sql").write_text(
        "{% macro passthrough(value) %}round({{ value }}, 2){% endmacro %}\n"
        "{% macro plus_one(value) %}{{ value }} + 1{% endmacro %}"
    )
    resource(root, [model(output={"sql": f"SELECT {expression} AS id FROM ref.items AS v", "columns": {"id": {}}})])
    result = optimize_project(project(root), preserve_marco=True)
    assert not result.skipped
    rewritten = project(root)
    saved = rewritten.models[0]
    tree = transformation_tree(saved.output, Compiler(rewritten.macros, preserve_macro=True), saved.name)
    assert any(call.name.lower() == "passthrough" for call in tree.find_all(exp.Anonymous))


@pytest.mark.parametrize(("body", "argument", "reason"), [
    ("SUM({{ value }})", "v.id", "aggregate/window/query"),
    ("{{ value }} + MAX(v.id)", "SUM(v.id)", "aggregate/window/query"),
    ("{{ value }} + {{ value }}", "SUM(v.id)", "aggregate/window/query"),
    ("SUM(v.id)", "SUM(v.other)", "aggregate/window/query"),
    ("0", "SUM(v.id)", "aggregate/window/query"),
    ("{{ value }} OVER ()", "SUM(v.id)", "aggregate/window/query"),
    ("(SELECT MAX(x.id) FROM ref.x AS x)", "v.id", "scalar SQL expression"),
    ("{{ value }} + v.hidden", "v.id", "column dependencies"),
])
def test_optimizer_rejects_macros_hiding_query_semantics_or_columns(root, body, argument, reason):
    (root / "macros/math.sql").write_text("{% macro hidden(value) %}" + body + "{% endmacro %}")
    resource(root, [model(output={"sql": f"SELECT hidden({argument}) AS id FROM ref.items AS v", "columns": {"id": {}}})])
    before = baseline(root)
    with pytest.raises(SpecError, match=reason):
        optimize_project(project(root), preserve_marco=True)
    assert baseline(root) == before


def test_optimizer_duckdb_semantics(root):
    import duckdb
    edit(root / "slbt_project.yaml", optimization={"dialect": "duckdb"})
    authored = "WITH x AS (SELECT 1 AS id UNION ALL SELECT 2 AS id) SELECT id + 0 AS id FROM x WHERE id > 0"
    resource(root, [model(output={"sql": authored, "columns": {"id": {}}})])
    optimize_project(project(root))
    compiled = dbt_outputs(project(root))["models/example.sql"].decode()
    from jinja2 import Environment
    executable = Environment().from_string(compiled).render(config=lambda value: "")
    with duckdb.connect() as connection:
        assert connection.execute(authored).fetchall() == connection.execute(executable).fetchall()


def test_cli_selection_boolean_values_discovery_and_init(root, monkeypatch):
    resource(root, [model(), model("second")])
    runner = CliRunner()
    for arguments in [
        ["validate"], ["--help"],
        ["optimize", "-s", "example", "second", "--preserve-marco", "true", "--enrich", "false", "--dry-run"],
        ["enrich", "-s", "example", "-s", "second", "--dry-run"],
        ["export", "catalog", "--catalog", "catalog", "--dry-run"],
    ]:
        result = runner.invoke(cli, ["--project-dir", str(root), *arguments])
        assert result.exit_code == 0, result.output
    assert runner.invoke(cli, ["--project-dir", str(root), "optimize", "-s"]).exit_code == 2
    assert runner.invoke(cli, ["--project-dir", str(root), "optimize", "--preserve-marco"]).exit_code == 2
    monkeypatch.chdir(root / "models")
    monkeypatch.delenv("slbt_PROJECT_DIR", raising=False)
    assert resolve_project() == root / "slbt_project.yaml"
    with pytest.raises(ValueError):
        initialize(root)


def test_python_project_is_operation_local_snapshot(root):
    loaded = project(root)
    resource(root, [model(description="Changed after load")])
    assert "description" not in loaded.models[0].authored
    assert project(root).models[0].authored["description"] == "Changed after load"
